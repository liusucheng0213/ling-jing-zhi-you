"""管理后台-游客反馈与关注点分析API。"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from html import escape
from io import BytesIO
from typing import Any
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from sqlalchemy import desc, func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models import BehaviorLog, ChatMessage, EmotionHistory, Session as ChatSession
from app.utils.jwt import get_current_admin

router = APIRouter(tags=["管理后台-分析"])

MONTHLY_REPORT_EVENT_TYPE = "analytics_monthly_report"
CHINA_STANDARD_TIME = timezone(timedelta(hours=8), name="CST")


def _report_now() -> datetime:
    """返回报告面向展示人员使用的中国标准时间。"""

    return datetime.now(CHINA_STANDARD_TIME)


TOPIC_KEYWORDS = {
    "门票": ["门票", "票价", "多少钱", "收费", "免费"],
    "开放时间": ["开放时间", "几点", "闭园", "开门", "营业"],
    "交通": ["怎么去", "交通", "公交", "地铁", "停车", "路线"],
    "餐饮": ["吃", "餐饮", "美食", "素斋", "饭店", "餐厅"],
    "表演": ["表演", "演出", "梵宫", "时间", "节目"],
    "路线": ["路线", "怎么走", "行程", "攻略", "推荐"],
    "无障碍": ["无障碍", "轮椅", "老人", "适老", "婴儿车"],
    "拍照点": ["拍照", "摄影", "打卡", "出片"],
    "历史文化": ["历史", "文化", "典故", "佛教", "传说", "故事"],
    "投诉建议": ["不好", "差", "投诉", "建议", "排队", "太贵", "不方便"],
}

NEGATIVE_WORDS = ["不好", "差", "失望", "投诉", "排队", "太贵", "不方便", "慢", "找不到"]
POSITIVE_WORDS = ["好", "棒", "喜欢", "满意", "漂亮", "方便", "不错", "开心"]


def _report_docx_bytes(
    payload: dict[str, Any],
    days: int,
    title: str = "灵境智游游客感受度分析报告",
    period_label: str | None = None,
) -> bytes:
    """生成可下载、可归档的标准DOCX报告。"""

    from docx import Document

    document = Document()
    document.add_heading(title, 0)
    document.add_paragraph(
        f"统计周期：{period_label or f'近 {days} 天'}；"
        f"数据类型：{'样例数据' if payload.get('is_sample') else '真实数据'}；"
        f"生成时间：{_report_now().strftime('%Y-%m-%d %H:%M')}"
    )
    overview = payload["overview"]
    document.add_heading("核心指标", level=1)
    document.add_paragraph(
        f"服务/问答量：{overview.get('total_qa', 0)}；平均满意度：{overview.get('avg_satisfaction') or '暂无评价'}；"
        f"正面率：{overview.get('positive_rate', 0)}%；负面率：{overview.get('negative_rate', 0)}%；"
        f"反馈数：{overview.get('feedback_count', 0)}。"
    )
    document.add_heading("游客反馈摘要", level=1)
    document.add_paragraph(payload["feedback"].get("summary", "暂无反馈摘要。"))
    document.add_heading("游客关注点", level=1)
    for concern in payload.get("concerns", []):
        document.add_paragraph(
            f"{concern.get('topic')}：{concern.get('mention_count')}次，{concern.get('sentiment')}",
            style="List Bullet",
        )
    document.add_heading("服务改进建议", level=1)
    for suggestion in payload.get("suggestions", []):
        document.add_paragraph(suggestion, style="List Number")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _report_pdf_bytes(payload: dict[str, Any], days: int) -> bytes:
    """生成支持中文、自动换行和跨页表格的PDF分析报告。"""

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape as xml_escape

    font_name = "STSong-Light"
    pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=18 * mm,
        bottomMargin=17 * mm,
        title="灵境智游游客感受度分析报告",
        author="灵境智游",
    )
    base_styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ChineseTitle",
        parent=base_styles["Title"],
        fontName=font_name,
        fontSize=20,
        leading=28,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#14532d"),
        spaceAfter=8,
        wordWrap="CJK",
    )
    section_style = ParagraphStyle(
        "ChineseSection",
        parent=base_styles["Heading2"],
        fontName=font_name,
        fontSize=13,
        leading=19,
        textColor=colors.HexColor("#166534"),
        spaceBefore=10,
        spaceAfter=6,
        keepWithNext=True,
        wordWrap="CJK",
    )
    body_style = ParagraphStyle(
        "ChineseBody",
        parent=base_styles["BodyText"],
        fontName=font_name,
        fontSize=9.5,
        leading=15,
        textColor=colors.HexColor("#1f2937"),
        wordWrap="CJK",
    )
    small_style = ParagraphStyle(
        "ChineseSmall",
        parent=body_style,
        fontSize=8,
        leading=12,
    )
    note_style = ParagraphStyle(
        "ChineseNote",
        parent=body_style,
        fontSize=8.5,
        leading=13,
        textColor=colors.HexColor("#92400e"),
        borderColor=colors.HexColor("#fde68a"),
        borderWidth=0.6,
        borderPadding=7,
        backColor=colors.HexColor("#fffbeb"),
        spaceAfter=8,
    )

    def paragraph(value: Any, style: ParagraphStyle = body_style) -> Paragraph:
        text_value = xml_escape(str(value if value is not None else "-"))
        return Paragraph(text_value.replace("\n", "<br/>"), style)

    common_table_style = TableStyle(
        [
            ("FONTNAME", (0, 0), (-1, -1), font_name),
            ("FONTSIZE", (0, 0), (-1, 0), 8.5),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dcfce7")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#14532d")),
            ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ]
    )

    overview = payload.get("overview") or {}
    feedback = payload.get("feedback") or {}
    generated_at = _report_now().strftime("%Y-%m-%d %H:%M")
    data_label = "演示样例数据" if payload.get("is_sample") else "实际运行数据"
    story: list[Any] = [
        Paragraph("灵境智游游客感受度分析报告", title_style),
        Paragraph(f"统计周期：近 {days} 天　　生成时间：{generated_at}　　数据类型：{data_label}", small_style),
        Spacer(1, 5),
        Paragraph(
            "说明：本报告依据游客问答、显式评价与系统行为记录生成。演示样例仅用于本机展示界面和报表能力，不代表景区真实客流或经营结果。"
            if payload.get("is_sample")
            else "说明：本报告依据游客问答、显式评价与系统行为记录生成，用于运营复盘与知识库优化。",
            note_style,
        ),
        Paragraph("一、核心指标", section_style),
    ]

    kpi_table = Table(
        [
            [paragraph("服务会话", small_style), paragraph("问答量", small_style), paragraph("平均满意度", small_style), paragraph("正面评价率", small_style), paragraph("低分反馈", small_style)],
            [
                paragraph(overview.get("total_visitors", 0), body_style),
                paragraph(overview.get("total_qa", 0), body_style),
                paragraph(f"{overview.get('avg_satisfaction') or '暂无'} / 5", body_style),
                paragraph(f"{overview.get('positive_rate', 0)}%", body_style),
                paragraph(feedback.get("low_rating_count", 0), body_style),
            ],
        ],
        colWidths=[35 * mm] * 5,
        repeatRows=1,
    )
    kpi_table.setStyle(common_table_style)
    story.extend(
        [
            kpi_table,
            Paragraph("二、游客反馈摘要", section_style),
            paragraph(feedback.get("summary") or "暂无反馈摘要。"),
            Paragraph("三、游客关注点分析", section_style),
        ]
    )

    concern_rows = [[paragraph(item, small_style) for item in ("关注点", "提及次数", "情感倾向", "趋势")]]
    sentiment_labels = {"positive": "正向", "negative": "负向", "neutral": "中性"}
    trend_labels = {"up": "上升", "down": "下降", "stable": "平稳"}
    for item in payload.get("concerns") or []:
        concern_rows.append(
            [
                paragraph(item.get("topic"), small_style),
                paragraph(item.get("mention_count", 0), small_style),
                paragraph(sentiment_labels.get(item.get("sentiment"), item.get("sentiment", "中性")), small_style),
                paragraph(trend_labels.get(item.get("trend"), item.get("trend", "平稳")), small_style),
            ]
        )
    if len(concern_rows) == 1:
        concern_rows.append([paragraph("暂无", small_style), paragraph(0, small_style), paragraph("-", small_style), paragraph("-", small_style)])
    concern_table = Table(concern_rows, colWidths=[58 * mm, 34 * mm, 42 * mm, 41 * mm], repeatRows=1)
    concern_table.setStyle(common_table_style)
    story.extend([concern_table, Paragraph("四、每日服务与情感趋势", section_style)])

    daily_rows = [[paragraph(item, small_style) for item in ("日期", "会话", "平均时长", "满意度", "正面率", "主要关注点")]]
    for item in payload.get("daily") or []:
        daily_rows.append(
            [
                paragraph(item.get("date"), small_style),
                paragraph(item.get("total_visitors", 0), small_style),
                paragraph(f"{item.get('avg_session_time', 0)} 分钟", small_style),
                paragraph(item.get("satisfaction") or "暂无", small_style),
                paragraph(f"{item.get('positive_rate', 0)}%", small_style),
                paragraph(item.get("top_concern", "-"), small_style),
            ]
        )
    if len(daily_rows) == 1:
        daily_rows.append([paragraph("暂无", small_style)] * 6)
    daily_table = Table(daily_rows, colWidths=[29 * mm, 22 * mm, 30 * mm, 28 * mm, 27 * mm, 39 * mm], repeatRows=1)
    daily_table.setStyle(common_table_style)
    story.extend([daily_table, Paragraph("五、分维度体验评分", section_style)])

    dimension_rows = [[paragraph(item, small_style) for item in ("体验维度", "平均分", "样本数")]]
    dimension_labels = {
        "answer_accuracy": "回答准确性",
        "guide_vividness": "讲解生动性",
        "avatar_voice": "数字人音色",
        "route_usefulness": "路线实用性",
    }
    for item in feedback.get("dimension_scores") or []:
        dimension_rows.append(
            [
                paragraph(dimension_labels.get(item.get("name"), item.get("name")), small_style),
                paragraph(f"{item.get('avg', 0)} / 5", small_style),
                paragraph(item.get("count", 0), small_style),
            ]
        )
    if len(dimension_rows) == 1:
        dimension_rows.append([paragraph("暂无分维度评分", small_style), paragraph("-", small_style), paragraph(0, small_style)])
    dimension_table = Table(dimension_rows, colWidths=[82 * mm, 46 * mm, 47 * mm], repeatRows=1)
    dimension_table.setStyle(common_table_style)
    story.extend([dimension_table, Paragraph("六、近期低分反馈", section_style)])

    low_feedback = feedback.get("recent_low_feedback") or []
    if low_feedback:
        for index, item in enumerate(low_feedback[:12], start=1):
            story.append(
                paragraph(
                    f"{index}. {str(item.get('created_at') or '-')[:10]}　{item.get('rating', '-')} 分　{item.get('feedback_text') or '未填写文字反馈'}"
                )
            )
    else:
        story.append(paragraph("暂无低分文字反馈。"))

    story.append(Paragraph("七、服务改进建议", section_style))
    suggestions = payload.get("suggestions") or ["继续积累游客互动数据，完善知识库覆盖。"]
    for index, suggestion in enumerate(suggestions, start=1):
        story.append(KeepTogether([paragraph(f"{index}. {suggestion}"), Spacer(1, 2)]))

    def draw_page(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont(font_name, 7.5)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(16 * mm, 9 * mm, "灵境智游 · 游客感受度分析")
        canvas.drawRightString(A4[0] - 16 * mm, 9 * mm, f"第 {doc.page} 页")
        canvas.restoreState()

    document.build(story, onFirstPage=draw_page, onLaterPages=draw_page)
    return output.getvalue()


def _monthly_archive_response(row: BehaviorLog) -> dict[str, Any]:
    """将月报行为日志转换为稳定的管理端列表结构。"""

    data = row.event_data or {}
    snapshot = data.get("snapshot") or {}
    overview = snapshot.get("overview") or {}
    return {
        "id": str(row.id),
        "label": data.get("label") or "近30天月报快照",
        "period_days": int(data.get("period_days") or 30),
        "period_start": data.get("period_start"),
        "period_end": data.get("period_end"),
        "generated_at": data.get("generated_at") or (row.created_at.isoformat() if row.created_at else None),
        "generated_by": data.get("generated_by"),
        "archive_mode": data.get("archive_mode") or "manual",
        "is_sample": bool(snapshot.get("is_sample")),
        "total_qa": int(overview.get("total_qa") or 0),
        "feedback_count": int(overview.get("feedback_count") or 0),
        "avg_satisfaction": overview.get("avg_satisfaction"),
    }


def _report_html(payload: dict[str, Any], days: int) -> str:
    """生成Word可打开的游客感受度报告HTML。"""

    overview = payload["overview"]
    feedback = payload["feedback"]
    generated_at = _report_now().strftime("%Y-%m-%d %H:%M")
    sample_note = "（样例数据）" if payload.get("is_sample") else "（真实数据）"

    daily_rows = "".join(
        "<tr>"
        f"<td>{escape(item.get('date', '-'))}</td>"
        f"<td>{item.get('total_visitors', 0)}</td>"
        f"<td>{item.get('satisfaction', 0)}</td>"
        f"<td>{item.get('positive_rate', 0)}%</td>"
        f"<td>{item.get('negative_rate', 0)}%</td>"
        f"<td>{escape(item.get('top_concern', '-'))}</td>"
        "</tr>"
        for item in payload.get("daily", [])
    )
    concern_rows = "".join(
        "<tr>"
        f"<td>{escape(item.get('topic', '-'))}</td>"
        f"<td>{item.get('mention_count', 0)}</td>"
        f"<td>{escape(item.get('sentiment', 'neutral'))}</td>"
        f"<td>{escape(item.get('trend', 'stable'))}</td>"
        "</tr>"
        for item in payload.get("concerns", [])
    )
    suggestion_items = "".join(f"<li>{escape(item)}</li>" for item in payload.get("suggestions", []))
    low_feedback_items = "".join(
        "<li>"
        f"{escape(str(item.get('created_at') or '-'))}，评分{item.get('rating', '-')}："
        f"{escape(str(item.get('feedback_text') or '未填写文字反馈'))}"
        "</li>"
        for item in feedback.get("recent_low_feedback", [])
    ) or "<li>暂无低分文字反馈。</li>"
    dimension_rows = "".join(
        "<tr>"
        f"<td>{escape(item.get('name', '-'))}</td>"
        f"<td>{item.get('avg', 0)} / 5</td>"
        f"<td>{item.get('count', 0)}</td>"
        "</tr>"
        for item in feedback.get("dimension_scores", [])
    )

    return f"""
<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <style>
    body {{ font-family: "Microsoft YaHei", Arial, sans-serif; color: #1f2937; line-height: 1.7; }}
    h1 {{ color: #0f766e; font-size: 26px; }}
    h2 {{ color: #115e59; font-size: 18px; margin-top: 24px; }}
    table {{ border-collapse: collapse; width: 100%; margin: 12px 0; }}
    th, td {{ border: 1px solid #cbd5e1; padding: 8px 10px; font-size: 12px; }}
    th {{ background: #ecfeff; color: #0f172a; }}
    .kpi td {{ width: 25%; text-align: center; }}
    .note {{ color: #92400e; background: #fffbeb; padding: 8px 12px; }}
  </style>
</head>
<body>
  <h1>灵境智游游客感受度分析报告 {sample_note}</h1>
  <p>统计周期：近 {days} 天；生成时间：{generated_at}</p>
  <p class="note">说明：报告基于游客问答、显式评价和情感识别结果生成，用于运营复盘和知识库优化。</p>

  <h2>一、核心指标</h2>
  <table class="kpi">
    <tr>
      <th>服务/问答量</th><th>平均满意度</th><th>正面评价率</th><th>低分反馈数</th>
    </tr>
    <tr>
      <td>{overview.get("total_visitors", 0)}</td>
      <td>{overview.get("avg_satisfaction", 0)} / 5</td>
      <td>{overview.get("positive_rate", 0)}%</td>
      <td>{feedback.get("low_rating_count", 0)}</td>
    </tr>
  </table>

  <h2>二、游客反馈摘要</h2>
  <p>{escape(feedback.get("summary", "暂无反馈摘要。"))}</p>

  <h2>三、游客关注点分析</h2>
  <table>
    <tr><th>关注点</th><th>提及次数</th><th>情感倾向</th><th>趋势</th></tr>
    {concern_rows or '<tr><td colspan="4">暂无关注点数据</td></tr>'}
  </table>

  <h2>四、情感与满意度日报</h2>
  <table>
    <tr><th>日期</th><th>服务/问答</th><th>满意度</th><th>正面率</th><th>负面率</th><th>主要关注点</th></tr>
    {daily_rows or '<tr><td colspan="6">暂无日报数据</td></tr>'}
  </table>

  <h2>五、近期低分反馈</h2>
  <ul>{low_feedback_items}</ul>

  <h2>六、分维度体验评分</h2>
  <table>
    <tr><th>维度</th><th>平均分</th><th>样本数</th></tr>
    {dimension_rows or '<tr><td colspan="3">暂无分维度评分</td></tr>'}
  </table>

  <h2>七、服务改进建议</h2>
  <ol>{suggestion_items or '<li>继续积累游客互动数据，完善知识库覆盖。</li>'}</ol>
</body>
</html>
"""


def _topic_for_text(text: str) -> str:
    lowered = text.lower()
    for topic, words in TOPIC_KEYWORDS.items():
        if any(word.lower() in lowered for word in words):
            return topic
    return "其他"


def _sentiment_for_text(text: str, rating: int | None = None) -> str:
    if rating is not None:
        if rating <= 2:
            return "negative"
        if rating >= 4:
            return "positive"
    lowered = text.lower()
    if any(word.lower() in lowered for word in NEGATIVE_WORDS):
        return "negative"
    if any(word.lower() in lowered for word in POSITIVE_WORDS):
        return "positive"
    return "neutral"


def _empty_payload(days: int, period_label: str | None = None) -> dict[str, Any]:
    return {
        "is_sample": False,
        "overview": {
            "total_visitors": 0,
            "avg_satisfaction": None,
            "positive_rate": 0,
            "negative_rate": 0,
            "total_qa": 0,
            "feedback_count": 0,
            "has_feedback": False,
        },
        "daily": [],
        "concerns": [],
        "feedback": {
            "summary": f"{period_label or f'近 {days} 天'}暂无真实问答或反馈数据。",
            "low_rating_count": 0,
            "rating_distribution": {"1": 0, "2": 0, "3": 0, "4": 0, "5": 0},
            "recent_low_feedback": [],
            "dimension_scores": [],
        },
        "suggestions": [],
    }


async def _build_rule_suggestions(summary: str, concerns: list[dict[str, Any]], low_rating_count: int) -> list[str]:
    """基于真实统计规则生成建议，不在报表请求链路中调用大模型。"""

    base = []
    top_topics = [item["topic"] for item in concerns[:3]]
    if "交通" in top_topics:
        base.append("补充交通、停车、接驳车相关知识，并在游客端快捷问题中前置。")
    if "表演" in top_topics:
        base.append("将表演时间、场次变更和预约方式做成高频快答。")
    if "门票" in top_topics:
        base.append("完善门票价格、优惠政策、套票和购票入口说明。")
    if low_rating_count:
        base.append("对低分反馈进行人工复核，优先处理排队、价格和路线不清晰问题。")
    if not base:
        base.append("继续沉淀问答和反馈样本，保持知识库每周补充。")
    if len(base) < 3:
        base.append("按高频问题复核知识库命中率，并补齐可核验的官方来源。")
    if len(base) < 3:
        base.append("持续监测响应时延与低置信回答，按周复盘异常会话。")
    return base[:3]


async def _build_analytics_summary(
    *,
    days: int,
    db: AsyncSession,
    start: datetime | None = None,
    end_exclusive: datetime | None = None,
    period_label: str | None = None,
) -> dict[str, Any]:
    """按明确时间窗口汇总分析；自动月报使用闭开自然月区间。"""

    start = start or (datetime.now() - timedelta(days=days))
    user_messages_query = (
        select(ChatMessage.content, ChatMessage.created_at)
        .where(ChatMessage.role == "user")
        .where(ChatMessage.created_at >= start)
        .order_by(ChatMessage.created_at)
        .limit(2000)
    )
    if end_exclusive is not None:
        user_messages_query = user_messages_query.where(ChatMessage.created_at < end_exclusive)
    user_messages_result = await db.execute(user_messages_query)
    user_messages = user_messages_result.all()

    feedback_query = (
        select(BehaviorLog.event_data, BehaviorLog.created_at)
        .where(BehaviorLog.event_type == "feedback")
        .where(BehaviorLog.created_at >= start)
        .order_by(desc(BehaviorLog.created_at))
        .limit(1000)
    )
    if end_exclusive is not None:
        feedback_query = feedback_query.where(BehaviorLog.created_at < end_exclusive)
    feedback_result = await db.execute(feedback_query)
    feedback_logs = feedback_result.all()
    is_sample = any(bool((event_data or {}).get("demo_seed")) for event_data, _ in feedback_logs)

    session_query = (
        select(ChatSession.started_at, ChatSession.ended_at)
        .where(ChatSession.started_at >= start)
        .order_by(ChatSession.started_at)
        .limit(2000)
    )
    if end_exclusive is not None:
        session_query = session_query.where(ChatSession.started_at < end_exclusive)
    session_rows = (await db.execute(session_query)).all()
    session_bucket: dict[str, list[float]] = defaultdict(list)
    current_time = datetime.now()
    for started_at, ended_at in session_rows:
        if not started_at:
            continue
        effective_end = ended_at or min(current_time, started_at + timedelta(minutes=30))
        duration_minutes = max(1.0, min(180.0, (effective_end - started_at).total_seconds() / 60))
        session_bucket[started_at.strftime("%Y-%m-%d")].append(duration_minutes)

    if not user_messages and not feedback_logs:
        return _empty_payload(days, period_label)

    topic_counts: Counter[str] = Counter()
    topic_sentiments: dict[str, Counter[str]] = defaultdict(Counter)
    daily_bucket: dict[str, dict[str, Any]] = defaultdict(lambda: {
        "total_visitors": 0,
        "qa": 0,
        "ratings": [],
        "positive": 0,
        "negative": 0,
        "topics": Counter(),
    })

    for content, created_at in user_messages:
        topic = _topic_for_text(content)
        sentiment = _sentiment_for_text(content)
        date_key = created_at.strftime("%Y-%m-%d") if created_at else datetime.now().strftime("%Y-%m-%d")
        topic_counts[topic] += 1
        topic_sentiments[topic][sentiment] += 1
        daily_bucket[date_key]["qa"] += 1
        daily_bucket[date_key]["topics"][topic] += 1
        if sentiment == "positive":
            daily_bucket[date_key]["positive"] += 1
        elif sentiment == "negative":
            daily_bucket[date_key]["negative"] += 1

    ratings: list[int] = []
    recent_low_feedback = []
    feedback_topics: Counter[str] = Counter()
    rating_distribution = {str(score): 0 for score in range(1, 6)}
    dimension_bucket: dict[str, list[int]] = defaultdict(list)
    for event_data, created_at in feedback_logs:
        data = event_data or {}
        rating = int(data.get("rating") or 0)
        text = str(data.get("feedback_text") or "")
        if rating:
            ratings.append(rating)
            rating_distribution[str(rating)] = rating_distribution.get(str(rating), 0) + 1
        if text:
            topic = _topic_for_text(text)
            sentiment = _sentiment_for_text(text, rating)
            topic_counts[topic] += 1
            topic_sentiments[topic][sentiment] += 1
            feedback_topics[topic] += 1
        if rating and rating <= 2:
            recent_low_feedback.append({
                "created_at": created_at.isoformat() if created_at else None,
                "rating": rating,
                "feedback_text": text,
                "tags": data.get("tags") or [],
            })
        for name, value in (data.get("dimension_scores") or {}).items():
            try:
                score = int(value)
            except (TypeError, ValueError):
                continue
            if 1 <= score <= 5:
                dimension_bucket[str(name)].append(score)

    concerns = []
    for topic, count in topic_counts.most_common(10):
        sentiments = topic_sentiments[topic]
        sentiment = sentiments.most_common(1)[0][0] if sentiments else "neutral"
        concerns.append({
            "topic": topic,
            "mention_count": count,
            "sentiment": sentiment,
            "trend": "up" if count >= 3 else "stable",
        })

    total_feedback = len(feedback_logs)
    avg_satisfaction = round(sum(ratings) / len(ratings), 2) if ratings else None
    positive_count = sum(1 for rating in ratings if rating >= 4)
    negative_count = sum(1 for rating in ratings if rating <= 2)
    positive_rate = round(positive_count / len(ratings) * 100) if ratings else 0
    negative_rate = round(negative_count / len(ratings) * 100) if ratings else 0

    daily = []
    for date_key in sorted(daily_bucket):
        bucket = daily_bucket[date_key]
        session_durations = session_bucket.get(date_key, [])
        top_topic = bucket["topics"].most_common(1)[0][0] if bucket["topics"] else "其他"
        total_emotion = bucket["positive"] + bucket["negative"]
        daily.append({
            "date": date_key,
            "total_visitors": len(session_durations) or bucket["qa"],
            "satisfaction": avg_satisfaction,
            "positive_rate": round(bucket["positive"] / total_emotion * 100) if total_emotion else positive_rate,
            "negative_rate": round(bucket["negative"] / total_emotion * 100) if total_emotion else negative_rate,
            "top_concern": top_topic,
            "avg_session_time": round(sum(session_durations) / len(session_durations), 1) if session_durations else 0,
        })

    hot_feedback_topic = feedback_topics.most_common(1)[0][0] if feedback_topics else (concerns[0]["topic"] if concerns else "暂无")
    summary_period = period_label or f"近{days}天"
    if total_feedback:
        summary = (
            f"{summary_period}共收到{total_feedback}条反馈，平均满意度{avg_satisfaction}/5，"
            f"低分反馈{negative_count}条；游客主要关注{hot_feedback_topic}。"
        )
    else:
        summary = (
            f"{summary_period}有{len(user_messages)}条真实游客提问，尚未收到满意度评价；"
            f"问答关注点主要为{hot_feedback_topic}。"
        )
    suggestions = await _build_rule_suggestions(summary, concerns, negative_count)

    return {
        "is_sample": is_sample,
        "overview": {
            "total_visitors": len(session_rows) or len(user_messages),
            "avg_satisfaction": avg_satisfaction,
            "positive_rate": positive_rate,
            "negative_rate": negative_rate,
            "total_qa": len(user_messages),
            "feedback_count": total_feedback,
            "has_feedback": bool(total_feedback),
        },
        "daily": daily,
        "concerns": concerns,
        "feedback": {
            "summary": summary,
            "low_rating_count": negative_count,
            "rating_distribution": rating_distribution,
            "recent_low_feedback": recent_low_feedback[:10],
            "dimension_scores": [
                {"name": name, "avg": round(sum(values) / len(values), 2), "count": len(values)}
                for name, values in dimension_bucket.items()
            ],
        },
        "suggestions": suggestions,
    }


@router.get("/admin/analytics/summary", summary="反馈、关注点与服务建议汇总")
async def get_analytics_summary(
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    admin: dict = Depends(get_current_admin),
):
    return await _build_analytics_summary(days=days, db=db)


def previous_calendar_month_window(current: datetime) -> tuple[datetime, datetime]:
    """返回上一个完整自然月的闭开区间 [月初, 本月月初)。"""

    current_month_start = current.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    previous_month_last_day = current_month_start - timedelta(days=1)
    previous_month_start = previous_month_last_day.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    return previous_month_start, current_month_start


async def archive_monthly_report_snapshot(
    db: AsyncSession,
    *,
    generated_by: str,
    user_id: UUID | None = None,
    archive_mode: str = "manual",
    automation_key: str | None = None,
    now: datetime | None = None,
) -> BehaviorLog:
    """生成月报快照；手动为近30天，自动为上一个完整自然月。"""

    current = now or datetime.now()
    if archive_mode == "automatic":
        range_start, range_end_exclusive = previous_calendar_month_window(current)
        days = (range_end_exclusive.date() - range_start.date()).days
        period_start = range_start.date().isoformat()
        period_end = (range_end_exclusive - timedelta(days=1)).date().isoformat()
        month_label = range_start.strftime("%Y年%m月")
        label = f"{month_label}自动月报（{period_start} 至 {period_end}）"
        payload = await _build_analytics_summary(
            days=days,
            db=db,
            start=range_start,
            end_exclusive=range_end_exclusive,
            period_label=month_label,
        )
    else:
        days = 30
        period_start = (current - timedelta(days=days)).date().isoformat()
        period_end = current.date().isoformat()
        label = f"{period_start} 至 {period_end} 月报快照"
        payload = await _build_analytics_summary(
            days=days,
            db=db,
            start=current - timedelta(days=days),
        )
    row = BehaviorLog(
        user_id=user_id,
        event_type=MONTHLY_REPORT_EVENT_TYPE,
        event_data={
            "schema_version": 2 if archive_mode == "automatic" else 1,
            "label": label,
            "period_days": days,
            "period_start": period_start,
            "period_end": period_end,
            "generated_at": current.isoformat(),
            "generated_by": generated_by,
            "archive_mode": archive_mode,
            "automation_key": automation_key,
            "snapshot": payload,
        },
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return row


async def ensure_automatic_monthly_report(
    db: AsyncSession,
    now: datetime | None = None,
) -> BehaviorLog | None:
    """在数据库事务锁内幂等归档上一个完整自然月。"""

    current = now or datetime.now()
    period_start, _ = previous_calendar_month_window(current)
    automation_key = f"completed-calendar-month:v2:{period_start.strftime('%Y-%m')}"

    # 多 worker 共享 PostgreSQL 事务级 advisory lock。持锁后再查再写，
    # 保证没有表结构迁移时同一完成月份也只生成一份自动归档。
    await db.execute(
        text(
            "SELECT pg_advisory_xact_lock("
            "hashtext(:lock_namespace), hashtext(:automation_key)"
            ")"
        ),
        {
            "lock_namespace": "lingjing.analytics.monthly-report.v2",
            "automation_key": automation_key,
        },
    )
    result = await db.execute(
        select(BehaviorLog.id)
        .where(BehaviorLog.event_type == MONTHLY_REPORT_EVENT_TYPE)
        .where(BehaviorLog.event_data["automation_key"].as_string() == automation_key)
        .limit(1)
    )
    if result.scalar_one_or_none() is not None:
        await db.commit()
        return None

    try:
        return await archive_monthly_report_snapshot(
            db,
            generated_by="system-monthly-scheduler",
            archive_mode="automatic",
            automation_key=automation_key,
            now=current,
        )
    except Exception:
        await db.rollback()
        raise


@router.get("/admin/analytics/monthly-reports", summary="获取已归档的月报快照")
async def list_monthly_reports(
    limit: int = Query(12, ge=1, le=36),
    db: AsyncSession = Depends(get_db),
    admin: dict = Depends(get_current_admin),
):
    """列出手动近30天快照和系统自动自然月归档。"""

    result = await db.execute(
        select(BehaviorLog)
        .where(BehaviorLog.event_type == MONTHLY_REPORT_EVENT_TYPE)
        .order_by(desc(BehaviorLog.created_at))
        .limit(limit)
    )
    return {"items": [_monthly_archive_response(row) for row in result.scalars().all()]}


@router.post("/admin/analytics/monthly-reports", status_code=status.HTTP_201_CREATED, summary="生成并归档近30天月报")
async def create_monthly_report(
    db: AsyncSession = Depends(get_db),
    admin: dict = Depends(get_current_admin),
):
    """人工生成临时近30天快照；自动月度归档由后台调度器负责。"""

    row = await archive_monthly_report_snapshot(
        db,
        generated_by=admin.get("phone") or admin.get("user_id") or "admin",
        user_id=UUID(admin["user_id"]) if admin.get("user_id") else None,
        archive_mode="manual",
    )
    return _monthly_archive_response(row)


@router.get("/admin/analytics/monthly-reports/{report_id}/report.docx", summary="下载已归档月报DOCX")
async def export_archived_monthly_report_docx(
    report_id: UUID,
    db: AsyncSession = Depends(get_db),
    admin: dict = Depends(get_current_admin),
):
    """从数据库快照导出DOCX，避免历史归档随实时数据变化。"""

    result = await db.execute(
        select(BehaviorLog)
        .where(BehaviorLog.id == report_id)
        .where(BehaviorLog.event_type == MONTHLY_REPORT_EVENT_TYPE)
    )
    row = result.scalar_one_or_none()
    if row is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="月报归档不存在")
    data = row.event_data or {}
    payload = data.get("snapshot")
    if not isinstance(payload, dict):
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="月报归档数据不完整")
    days = int(data.get("period_days") or 30)
    if data.get("archive_mode") == "automatic":
        period_start = str(data.get("period_start") or "")
        month_title = period_start[:7].replace("-", "年") + "月" if len(period_start) >= 7 else "自然月"
        title = f"灵境智游{month_title}运营月报"
        period_label = f"{data.get('period_start')} 至 {data.get('period_end')}"
    else:
        title = "灵境智游近30天运营月报"
        period_label = None
    report_bytes = _report_docx_bytes(payload, days, title=title, period_label=period_label)
    filename = f"lingjing-monthly-report-{str(report_id)[:8]}.docx"
    return Response(
        content=report_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/admin/analytics/report.doc", summary="导出游客感受度Word报告")
async def export_analytics_report_doc(
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    admin: dict = Depends(get_current_admin),
):
    """导出Word兼容的游客关注点、情感趋势与服务建议报告。"""

    payload = await get_analytics_summary(days=days, db=db, admin=admin)
    html = _report_html(payload, days)
    filename = f"lingjing-tourist-sentiment-report-{_report_now().strftime('%Y%m%d')}.doc"
    return Response(
        content=html.encode("utf-8"),
        media_type="application/msword; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/admin/analytics/report.docx", summary="导出游客感受度DOCX报告")
async def export_analytics_report_docx(
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    admin: dict = Depends(get_current_admin),
):
    """导出标准DOCX游客感受度报告。"""

    payload = await get_analytics_summary(days=days, db=db, admin=admin)
    filename = f"lingjing-tourist-sentiment-report-{_report_now().strftime('%Y%m%d')}.docx"
    return Response(
        content=_report_docx_bytes(payload, days),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/admin/analytics/report.pdf", summary="导出游客感受度PDF报告")
async def export_analytics_report_pdf(
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    admin: dict = Depends(get_current_admin),
):
    """导出PDF游客感受度报告。"""

    payload = await get_analytics_summary(days=days, db=db, admin=admin)
    filename = f"lingjing-tourist-sentiment-report-{_report_now().strftime('%Y%m%d')}.pdf"
    return Response(
        content=_report_pdf_bytes(payload, days),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/admin/analytics/emotion-trend", summary="真实情感趋势")
async def get_real_emotion_trend(
    days: int = Query(7, ge=1, le=90),
    db: AsyncSession = Depends(get_db),
    admin: dict = Depends(get_current_admin),
):
    start = datetime.now() - timedelta(days=days)
    result = await db.execute(
        select(
            func.date(EmotionHistory.created_at).label("date"),
            EmotionHistory.emotion_type,
            func.count(EmotionHistory.id).label("count"),
        )
        .where(EmotionHistory.created_at >= start)
        .group_by("date", EmotionHistory.emotion_type)
        .order_by("date")
    )
    rows = result.all()
    if not rows:
        return {"is_sample": False, "items": []}
    grouped: dict[str, dict[str, int]] = defaultdict(dict)
    for date, emotion, count in rows:
        key = date.strftime("%Y-%m-%d") if hasattr(date, "strftime") else str(date)
        grouped[key][emotion] = count
    items = []
    for date, emotions in sorted(grouped.items()):
        total = sum(emotions.values()) or 1
        items.append({"date": date, **{key: round(value / total, 3) for key, value in emotions.items()}})
    return {"is_sample": False, "items": items}
