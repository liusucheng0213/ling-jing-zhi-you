#!/usr/bin/env python3
"""同步并规范化 2026 软件杯 A5 官方公开资料包。

默认只读取显式传入的本地 ZIP；加 ``--download`` 时才从官网获取原包。
原始 ZIP 不进入 Git，生成的 JSON 可直接随仓库完成干净环境初始化。
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import urllib.request
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from docx import Document


A5_PAGE_URL = "https://www.cnsoftbei.com/content-3-1245-1.html"
A5_PACKAGE_URL = "https://www.cnsoftbei.com/uploadfile/2026/0323/20260323113204906.zip"
A5_PACKAGE_SHA256 = "833c57758d9daee7394b91efde7de035422f8c04c4529025a78950b29411889c"
A5_PUBLISHED_DATE = "2026-03-23"
A5_PACKAGE_NAME = "2026软件杯A5示范景区公开资料包"
OUTPUT_DIR = Path(__file__).resolve().parent
EXPECTED_SOURCE_HASHES = {
    "景点景区旅游数据行为分析数据.xlsx": "173113373ecb75a2e6e572f3a6396597bc66be1e5af68547d04002d36251af52",
    "灵山胜境 景点结构化数据集.docx": "f320d4677a6e6af72f754180f7a5f304519860a08768d94719901b2f9c9af99c",
    "灵山胜境：历史、文化、景点特色与个性化游览指南.docx": "88ce8d303b9b65c5f8764db2f1bbc06dd7e4bfcc5381dca707b2409a05782a07",
}

FIELD_QUESTIONS = (
    ("具体位置", "{name}具体位于哪里？"),
    ("建筑/景观参数", "{name}有哪些可核验的建筑或景观参数？"),
    ("文化内涵", "{name}的文化内涵是什么？"),
    ("游玩亮点", "{name}最值得看的游玩亮点是什么？"),
)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _read_zip(path: Path | None, download: bool) -> bytes:
    if path:
        return path.read_bytes()
    if not download:
        raise SystemExit("请通过 --zip 指定官网资料包，或显式使用 --download。")
    request = urllib.request.Request(A5_PACKAGE_URL, headers={"User-Agent": "Lingjing-A5-Sync/1.0"})
    with urllib.request.urlopen(request, timeout=60) as response:
        return response.read()


def _document_text(document: Document) -> str:
    sections = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip().replace("\n", "；") for cell in row.cells]
            if any(values):
                sections.append(" | ".join(values))
    return "\n\n".join(sections).strip()


def _structured_rows(document: Document) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for table in document.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if "景点ID" not in headers or "景点名称" not in headers:
            continue
        for row in table.rows[1:]:
            values = [cell.text.strip().replace("\n", "；") for cell in row.cells]
            record = dict(zip(headers, values))
            if record.get("景点名称"):
                rows.append(record)
    return rows


def _known_conflicts(spot_id: str) -> list[dict[str, Any]]:
    """记录资料包自身可复核的冲突，不替评委擅自裁决。"""

    if spot_id != "LS-006":
        return []
    return [
        {
            "field": "九龙灌浴总体参数",
            "status": "unresolved",
            "structured_value": "总高27.2m，整体耗铜180吨",
            "narrative_value": "叙述性指南出现总高27.5m、耗铜260余吨的不同口径",
            "handling": "展示差异，不把任一口径包装为已核实实时事实",
        }
    ]


def _normalize_entries(entries: list[tuple[str, bytes]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    documents: list[dict[str, Any]] = []
    source_files: list[dict[str, Any]] = []
    narrative: tuple[str, bytes] | None = None

    for file_name, payload in entries:
        source_files.append(
            {
                "source_file": file_name,
                "sha256": _sha256(payload),
                "bytes": len(payload),
                "role": "behavior_sample" if file_name.lower().endswith(".xlsx") else "knowledge_source",
            }
        )
        if not file_name.lower().endswith(".docx"):
            continue
        document = Document(io.BytesIO(payload))
        rows = _structured_rows(document)
        if not rows:
            narrative = (file_name, payload)
            continue
        for record in rows:
            scenic_area = record.get("景区名称") or "灵山胜境"
            source_spot_id = record.get("景点ID") or None
            conflicts = _known_conflicts(source_spot_id or "")
            content = "\n".join(
                f"{field}：{record.get(field, '')}"
                for field in record
                if record.get(field)
            )
            documents.append(
                {
                    "title": f"A5公开资料-{scenic_area}-{record['景点名称']}",
                    "doc_type": "spot_reference",
                    # 拈花湾不创建 Spot/路线/GPS 外键，仅在 metadata 保留官方编号。
                    "spot_id": source_spot_id if scenic_area == "灵山胜境" else None,
                    "content": content,
                    "source_file": file_name,
                    "metadata": {
                        "source_package": A5_PACKAGE_NAME,
                        "source_package_url": A5_PACKAGE_URL,
                        "source_package_hash": A5_PACKAGE_SHA256,
                        "source_file": file_name,
                        "source_hash": _sha256(payload),
                        "source_spot_id": source_spot_id,
                        "scenic_area": scenic_area,
                        "validity": "stable_with_historical_dynamic_fields",
                        "conflict_status": "unresolved" if conflicts else "none",
                        "conflicts": conflicts,
                    },
                }
            )

    if narrative:
        file_name, payload = narrative
        documents.append(
            {
                "title": "A5公开资料-灵山胜境历史文化与个性化游览指南",
                "doc_type": "official_package_guide",
                "spot_id": None,
                "content": _document_text(Document(io.BytesIO(payload))),
                "source_file": file_name,
                "metadata": {
                    "source_package": A5_PACKAGE_NAME,
                    "source_package_url": A5_PACKAGE_URL,
                    "source_package_hash": A5_PACKAGE_SHA256,
                    "source_file": file_name,
                    "source_hash": _sha256(payload),
                    "scenic_area": "灵山胜境",
                    "validity": "mixed_historical_baseline",
                    "conflict_status": "contains_unresolved_conflicts",
                    "conflicts": _known_conflicts("LS-006"),
                },
            }
        )

    area_counts: dict[str, int] = {}
    for item in documents:
        area = str(item["metadata"].get("scenic_area") or "未标注")
        area_counts[area] = area_counts.get(area, 0) + 1
    manifest = {
        "title": A5_PACKAGE_NAME,
        "problem_code": "A5",
        "problem_title": "大模型赋能的5A级景区AI数字人智能导览系统",
        "organizer": "锐捷网络股份有限公司苏州分公司",
        "official_page": A5_PAGE_URL,
        "official_package": A5_PACKAGE_URL,
        "published_date": A5_PUBLISHED_DATE,
        "package_sha256": A5_PACKAGE_SHA256,
        "scoring": {"功能完整度": 40, "技术创新": 30, "行业体验": 20, "文档": 10},
        "hard_metrics": {"事实问答准确率": ">=90%", "语音问答延迟": "<5s"},
        "scope": {
            "route_and_location": "仅灵山胜境",
            "knowledge": ["灵山胜境", "拈花湾禅意小镇"],
            "nianhua_bay": "仅知识问答，不创建路线、GPS点位或顶级页面",
        },
        "source_files": source_files,
        "normalized_documents": len(documents),
        "area_counts": area_counts,
        "conflicts": _known_conflicts("LS-006"),
    }
    return documents, manifest


def _entries_from_zip(zip_bytes: bytes) -> list[tuple[str, bytes]]:
    with ZipFile(io.BytesIO(zip_bytes), metadata_encoding="gbk") as archive:
        return [
            (Path(info.filename).name, archive.read(info))
            for info in archive.infolist()
            if not info.is_dir()
        ]


def _entries_from_directory(source_dir: Path) -> list[tuple[str, bytes]]:
    allowed = {".docx", ".xlsx"}
    entries = [
        (path.name, path.read_bytes())
        for path in sorted(source_dir.rglob("*"))
        if path.is_file() and path.suffix.lower() in allowed
    ]
    if not entries:
        raise SystemExit(f"资料目录中未找到 DOCX/XLSX: {source_dir}")
    actual_hashes = {name: _sha256(payload) for name, payload in entries}
    mismatches = {
        name: {"expected": expected, "actual": actual_hashes.get(name)}
        for name, expected in EXPECTED_SOURCE_HASHES.items()
        if actual_hashes.get(name) != expected
    }
    if mismatches:
        raise SystemExit(f"解压资料文件与已核验官网包不一致: {json.dumps(mismatches, ensure_ascii=False)}")
    return entries


def _field_map(document: dict[str, Any]) -> dict[str, str]:
    result: dict[str, str] = {}
    for line in str(document["content"]).splitlines():
        if "：" not in line:
            continue
        key, value = line.split("：", 1)
        result[key] = value.strip()
    return result


def _evidence_token(value: str) -> str:
    """选取回答中可确定复现的短证据词，不调用第二个模型评分。"""

    measure = re.search(r"\d+(?:\.\d+)?\s*(?:m|㎡|米|吨|座|组|只|年)", value, flags=re.I)
    if measure:
        return re.sub(r"\s+", "", measure.group(0))
    clause = re.split(r"[，,；;。]", value, maxsplit=1)[0].strip()
    clause = re.sub(r"^(位于|地处|核心为|代表|象征|以)", "", clause)
    return clause[:10] if len(clause) > 10 else clause


def _evidence_tokens(field: str, value: str) -> list[str]:
    """为同一事实提供少量可解释的等价短语，容许自然语言插入助词。"""

    primary = _evidence_token(value)
    tokens = [primary]
    if field == "具体位置":
        location_markers = (
            "入口",
            "中轴",
            "西侧",
            "东侧",
            "南侧",
            "北侧",
            "大佛座基",
            "香水海",
            "湖畔",
            "菩提大道",
            "核心区",
        )
        tokens.extend(marker for marker in location_markers if marker in value)
    return list(dict.fromkeys(token for token in tokens if token))


def _build_benchmark(documents: list[dict[str, Any]]) -> list[dict[str, Any]]:
    spot_docs = [item for item in documents if item["doc_type"] == "spot_reference"]
    cases: list[dict[str, Any]] = []

    def add_stable(items: list[dict[str, Any]], limit: int, prefix: str) -> None:
        before = len(cases)
        for document in items:
            fields = _field_map(document)
            name = fields.get("景点名称", "")
            for field, template in FIELD_QUESTIONS:
                value = fields.get(field, "")
                if not name or not value:
                    continue
                tokens = _evidence_tokens(field, value)
                cases.append(
                    {
                        "id": f"{prefix}-{len(cases) - before + 1:03d}",
                        "category": "stable_fact",
                        "scenic_area": document["metadata"]["scenic_area"],
                        "question": template.format(name=name),
                        "required_groups": [[name], tokens],
                        "forbidden": [],
                        "source_spot_id": document["metadata"].get("source_spot_id"),
                        "source_files": [document["source_file"]],
                    }
                )
                if len(cases) - before >= limit:
                    return

    lingshan = [item for item in spot_docs if item["metadata"]["scenic_area"] == "灵山胜境"]
    nianhua = [item for item in spot_docs if "拈花湾" in item["metadata"]["scenic_area"]]
    add_stable(lingshan, 60, "LS")
    add_stable(nianhua, 20, "NH")

    context_specs = (
        ("灵山大佛", "它的建筑参数是什么？", ["88"]),
        ("九龙灌浴", "它表现了什么文化故事？", ["释迦牟尼", "佛诞", "太子佛"]),
        ("灵山梵宫", "那里最值得看什么？", ["木雕", "星空穹顶", "华藏世界"]),
        ("五印坛城", "它主要展示哪一类文化？", ["藏传佛教"]),
        ("祥符禅寺", "它的历史可以追溯到什么时候？", ["唐", "贞观"]),
        ("阿育王柱", "它的建筑材质和造型呢？", ["花岗岩"]),
        ("佛足坛", "它有什么文化内涵？", ["佛足"]),
        ("灵山大照壁", "那里适合看什么？", ["照壁"]),
        ("香月花街", "那里有哪些游玩亮点？", ["花街", "青石板", "禅意商铺"]),
        ("拈花塔", "它的文化寓意是什么？", ["禅"]),
    )
    for index, (entity, follow_up, tokens) in enumerate(context_specs, start=1):
        area = "拈花湾禅意小镇" if entity in {"香月花街", "拈花塔"} else "灵山胜境"
        cases.append(
            {
                "id": f"CTX-{index:03d}",
                "category": "context_followup",
                "scenic_area": area,
                "turns": [f"请先简要介绍{entity}。", follow_up],
                "question": follow_up,
                # 多轮追问允许自然回答不机械复述景点名；正确实体由官方来源标题和
                # 景区范围共同验证，正文仍必须包含该实体的关键事实词。
                "required_groups": [tokens],
                "forbidden": [],
                "source_files": [],
            }
        )

    protection_specs = (
        ("今天九龙灌浴几点演出？", [["无法"], ["400-168-0303"]], ["10:35", "11:30"]),
        ("灵山胜境今天门票多少钱？", [["无法"], ["400-168-0303"]], ["210元", "105元"]),
        ("现在灵山梵宫开放吗？", [["无法"], ["现场公告", "400-168-0303"]], []),
        ("今日景区几点闭园？", [["无法"], ["现场公告", "400-168-0303"]], []),
        ("今天吉祥颂有哪几场？", [["无法"], ["现场公告", "400-168-0303"]], []),
        ("九龙灌浴总高到底是27.2米还是27.5米？", [["资料"], ["27.2"], ["27.5"]], []),
        ("九龙灌浴耗铜180吨和260吨哪个准确？", [["资料"], ["180"], ["260"]], []),
        ("资料里九龙灌浴参数为什么不一致？", [["冲突", "不同口径"], ["未核实", "无法确认"]], []),
        ("可以直接确认九龙灌浴总高27.5米吗？", [["不能", "无法"], ["27.2"], ["27.5"]], []),
        ("九龙灌浴260吨耗铜是确定事实吗？", [["不能", "无法"], ["180"], ["260"]], []),
    )
    for index, (question, required, forbidden) in enumerate(protection_specs, start=1):
        cases.append(
            {
                "id": f"SAFE-{index:03d}",
                "category": "dynamic_or_conflict",
                "scenic_area": "灵山胜境",
                "question": question,
                "required_groups": required,
                "forbidden": forbidden,
                "source_files": [],
                "must_pass": True,
            }
        )

    counts = {
        "lingshan": sum(1 for item in cases if item["id"].startswith("LS-")),
        "nianhua": sum(1 for item in cases if item["id"].startswith("NH-")),
        "context": sum(1 for item in cases if item["id"].startswith("CTX-")),
        "protection": sum(1 for item in cases if item["id"].startswith("SAFE-")),
    }
    if len(cases) != 100 or counts != {"lingshan": 60, "nianhua": 20, "context": 10, "protection": 10}:
        raise RuntimeError(f"A5 题集数量异常: total={len(cases)}, counts={counts}")
    return cases


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--zip", type=Path, help="官网下载的原始 ZIP 路径")
    parser.add_argument("--source-dir", type=Path, help="官网资料包的已解压目录")
    parser.add_argument("--download", action="store_true", help="从官网临时下载资料包")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()

    selected_sources = sum(bool(value) for value in (args.zip, args.source_dir, args.download))
    if selected_sources != 1:
        raise SystemExit("--zip、--source-dir 与 --download 必须且只能选择一种。")
    if args.source_dir:
        entries = _entries_from_directory(args.source_dir)
        source_mode = "verified_extracted_files"
    else:
        zip_bytes = _read_zip(args.zip, args.download)
        actual_hash = _sha256(zip_bytes)
        if actual_hash != A5_PACKAGE_SHA256:
            raise SystemExit(f"资料包 SHA-256 不匹配: expected={A5_PACKAGE_SHA256}, actual={actual_hash}")
        entries = _entries_from_zip(zip_bytes)
        source_mode = "verified_zip"
    documents, manifest = _normalize_entries(entries)
    manifest["normalization_source"] = source_mode
    benchmark = _build_benchmark(documents)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    outputs = {
        "a5_official_materials.json": documents,
        "a5_official_manifest.json": manifest,
        "a5_official_benchmark.json": benchmark,
    }
    for file_name, payload in outputs.items():
        target = args.output_dir / file_name
        target.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"已生成 {target} ({len(payload) if isinstance(payload, list) else 'manifest'})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
